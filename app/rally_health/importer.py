from __future__ import annotations

import hashlib
import json
import re
import unicodedata
from pathlib import Path
from typing import Any

import openpyxl

from .services.formula_registry import build_registry


def _text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _slug(value: str) -> str:
    text = unicodedata.normalize("NFKD", value.lower().strip())
    text = "".join(char for char in text if not unicodedata.combining(char))
    text = re.sub(r"[^a-z0-9]+", "_", text)
    return text.strip("_") or "exercise"


def _split_list(value: Any) -> list[str]:
    if value is None:
        return []
    return [item.strip() for item in re.split(r"[,;\n]", str(value)) if item.strip()]


def load_rally(path: Path) -> dict[str, Any]:
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    metrics: list[dict[str, Any]] = []
    for row_number, row in enumerate(workbook["Chi_so"].iter_rows(values_only=True), start=1):
        if row_number == 1 or not row or not row[0]:
            continue
        values = list(row) + [None] * max(0, 21 - len(row))
        metrics.append({
            "metric_code": _text(values[0]),
            "group": _text(values[1]),
            "name": _text(values[2]),
            "in_whoop_form": _text(values[3]),
            "evidence_level": _text(values[4]),
            "user_explanation": _text(values[5]),
            "reported_accuracy": _text(values[6]),
            "primary_citation": _text(values[7]),
            "paper_result": _text(values[8]),
            "other_citations": _text(values[9]),
            "recommendation": _text(values[10]),
            "calculation_description": _text(values[11]),
            "missing_input_factors": {
                name: values[index]
                for index, name in zip(range(12, 21), [
                    "missing_hrv", "missing_respiratory_rate", "missing_spo2",
                    "missing_sleep_stages", "sparse_heart_rate", "not_worn_during_sleep",
                    "iphone_time_in_bed_only", "basic_watch", "new_user_7_13_days",
                ])
            },
            "source": {"workbook": path.name, "sheet": "Chi_so", "row": row_number},
            "calculation_status": "review_required" if not values[11] else "description_only_not_executable",
        })

    supplemental: list[dict[str, Any]] = []
    sheet = workbook["Dan_chung_phu"]
    for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        if row_number == 1 or not row or not row[0]:
            continue
        values = list(row) + [None] * max(0, 6 - len(row))
        supplemental.append({
            "rule_id": _text(values[0]), "evidence_level": _text(values[1]),
            "construct": _text(values[2]), "notes": _text(values[3]),
            "citations_and_results": _text(values[4]), "recommendation": _text(values[5]),
            "source": {"workbook": path.name, "sheet": "Dan_chung_phu", "row": row_number},
            "status": "review_required",
        })
    return {"metrics": metrics, "supplemental_rules": supplemental}


def load_female_workouts(path: Path) -> list[dict[str, Any]]:
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheet = workbook["Bài tập"]
    records: list[dict[str, Any]] = []
    seen: set[str] = set()
    for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        if row_number <= 4 or not row or not row[0]:
            continue
        values = list(row) + [None] * max(0, 13 - len(row))
        name = _text(values[0]) or "exercise"
        signature = hashlib.sha256(f"female|{_slug(name)}".encode()).hexdigest()[:10]
        exercise_id = f"FE_{signature}"
        if exercise_id in seen:
            exercise_id = f"{exercise_id}_{row_number}"
        seen.add(exercise_id)
        records.append({
            "exercise_id": exercise_id, "id_generated": True, "sex_catalog": "female",
            "name": name, "video_label": _text(values[1]),
            "category_ids": _split_list(values[2]), "intensity": _text(values[3]),
            "description": _text(values[4]), "steps": _text(values[5]),
            "repetitions": _text(values[6]), "equipment": _split_list(values[7]),
            "kcal_per_minute_at_60kg": values[8] if isinstance(values[8], (int, float)) else None,
            "cycle_phases": _split_list(values[9]), "life_stages": _split_list(values[10]),
            "additional_information": _text(values[11]), "session_tags": _split_list(values[12]),
            "source": {"workbook": path.name, "sheet": "Bài tập", "row": row_number},
            "review_status": "needs_domain_review",
        })
    return records


def load_male_workouts(path: Path) -> list[dict[str, Any]]:
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=False)
    sheet = workbook["exercises"]
    records: list[dict[str, Any]] = []
    for row_number in range(6, sheet.max_row + 1):
        values = [sheet.cell(row=row_number, column=column).value for column in range(1, 20)]
        if not values[1] or not str(values[0] or "").strip().startswith(("A", "B")):
            continue
        records.append({
            "exercise_id": _text(values[1]), "id_generated": False, "sex_catalog": "male",
            "order": _text(values[0]), "name_vi": _text(values[2]), "name_en": _text(values[3]),
            "difficulty": values[4] if isinstance(values[4], (int, float)) else None,
            "reps_weak": _text(values[5]), "reps_medium": _text(values[6]),
            "reps_strong": _text(values[7]), "sets_weak": _text(values[8]),
            "sets_medium": _text(values[9]), "sets_strong": _text(values[10]),
            "recommended_for": _text(values[11]), "steps": _text(values[12]),
            "equipment": _split_list(values[13]),
            "met": values[14] if isinstance(values[14], (int, float)) else None,
            "evidence_source": _text(values[15]), "notes": _text(values[16]),
            "video_url": _text(values[17]), "video_timestamp": _text(values[18]),
            "source": {"workbook": path.name, "sheet": "exercises", "row": row_number},
            "review_status": "needs_domain_review",
        })
    return records


def load_roadmap(path: Path) -> list[dict[str, Any]]:
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheet = workbook["Content Roadmap"]
    records: list[dict[str, Any]] = []
    for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        if row_number == 1 or not row or not row[0]:
            continue
        values = list(row) + [None] * max(0, 10 - len(row))
        records.append({
            "source_id": _text(values[0]), "category": _text(values[1]),
            "title_en": _text(values[2]), "title_vi": _text(values[3]),
            "content_brief": _text(values[4]), "priority": _text(values[5]),
            "trigger_mapping": _text(values[6]), "source_url": _text(values[7]),
            "evidence_status": _text(values[8]), "content_status": _text(values[9]),
            "source": {"workbook": path.name, "sheet": "Content Roadmap", "row": row_number},
            "retrieval_eligible": False,
        })
    return records


def load_drkegel_metadata(path: Path) -> dict[str, Any]:
    return {
        "source_type": "competitor_research",
        "verified_medical_source": False,
        "recommendation_eligible": False,
        "knowledge_retrieval_eligible": False,
        "source": {"workbook": path.name},
        "warning": "Competitor material; claims and efficacy are not independently verified.",
    }


def import_workbooks(source_dir: Path, output_dir: Path) -> dict[str, Any]:
    source_dir = source_dir.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    rally_path = next(source_dir.glob("Rally*.xlsx"), None)
    female_path = source_dir / "BaiTap_Nu_GoogleSheets.xlsx"
    male_path = source_dir / "Ex_Male_Fitness_Huydq.xlsx"
    roadmap_path = source_dir / "Mens_Wellness_Content_Roadmap.xlsx"
    drkegel_path = source_dir / "DrKegel_BaiTap.xlsm"
    required = [p for p in [rally_path, female_path, male_path, roadmap_path, drkegel_path] if p is None or not p.exists()]
    if required:
        raise FileNotFoundError("Missing expected source workbook(s) in " + str(source_dir))

    rally = load_rally(rally_path)
    female = load_female_workouts(female_path)
    male = load_male_workouts(male_path)
    roadmap = load_roadmap(roadmap_path)
    drkegel = load_drkegel_metadata(drkegel_path)
    registry_metrics = build_registry(rally["metrics"])
    data = {
        "schema_version": "1.0.0",
        "source_directory": str(source_dir),
        "rally": {"metrics": registry_metrics, "supplemental_rules": rally["supplemental_rules"]},
        "metric_registry": {"version": "1.0.0", "metrics": registry_metrics},
        "workouts": {"female": female, "male": male},
        "mens_wellness_roadmap": roadmap,
        "drkegel_research_metadata": drkegel,
    }
    for name, payload in data.items():
        if name == "source_directory":
            continue
        (output_dir / f"{name}.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    report = {
        "schema_version": "1.0.0",
        "metrics_count": len(rally["metrics"]),
        "supplemental_rule_count": len(rally["supplemental_rules"]),
        "female_workout_count": len(female),
        "female_missing_session_tags": sum(not row["session_tags"] for row in female),
        "female_missing_cycle_phase": sum(not row["cycle_phases"] for row in female),
        "male_workout_count": len(male),
        "male_missing_video_url": sum(not row["video_url"] for row in male),
        "male_missing_video_timestamp": sum(not row["video_timestamp"] for row in male),
        "roadmap_topic_count": len(roadmap),
        "roadmap_not_retrieval_eligible": sum(not row["retrieval_eligible"] for row in roadmap),
        "rally_missing_calculation_descriptions": [row["metric_code"] for row in rally["metrics"] if not row["calculation_description"]],
        "rally_review_required_rules": len(rally["supplemental_rules"]),
        "metric_registry_version": "1.0.0",
        "metric_registry_status_counts": {
            status: sum(item["formula_registry"]["status"] == status for item in registry_metrics)
            for status in sorted({item["formula_registry"]["status"] for item in registry_metrics})
        },
        "metric_registry_enabled_aggregations": [
            item["metric_code"] for item in registry_metrics if item["formula_registry"]["executable"]
        ],
        "warnings": [
            "Enabled formulas currently perform raw-observation median aggregation only; they are not Rally-derived formulas.",
            "Rally calculation descriptions and supplemental rules require review before executable implementation.",
            "DrKegel and Planned roadmap content are excluded from health recommendations/retrieval.",
            "Generated female exercise IDs are provisional and must be checked for stable collisions before production.",
        ],
    }
    (output_dir / "import_report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return report
